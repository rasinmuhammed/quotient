"""Build the eval set: render each RFQ spec into the file a customer would send, and compute its gold answer.

    uv run python evals/build.py

Specs (evals/specs/*.yaml) hold what a person decided: the document text, and for each line the SKU it means,
the quantity and unit as written, and the flags a reviewer would raise that only a reader can know (no match,
grade not stated, a unit that cannot be converted). Every expected number is then computed by the pricing
engine from those gold lines; none is typed by hand. Flags that follow from the data (below MOQ, expired price,
margin floor, hazmat, no freight lane) come from the engine too, and a spec's `expect` list is asserted against
them so a data change cannot silently remove a trap.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, ConfigDict, field_validator
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from quotient.catalog import Catalog, default_catalog
from quotient.domain import SEVERITY, FlagCode, PricingRequest, Severity
from quotient.grades import GRADE_ORDER, default_grade
from quotient.pricing.engine import PricingEngine, quote_total
from quotient.pricing.rules import default_rules
from quotient.pricing.units import to_uom

HERE = Path(__file__).parent
IST = timezone(timedelta(hours=5, minutes=30))  # the sales desk's clock; received dates are read in it
SPECS, RFQS, GOLD = HERE / "specs", HERE / "rfqs", HERE / "gold"
SALES_INBOX = "Sales <sales@quotient-demo.example>"

# Chosen before any prompt existed, balanced so each trap type appears on both sides where it occurs twice.
HOLDOUT = ["rfq-004", "rfq-008", "rfq-012", "rfq-015", "rfq-017", "rfq-019", "rfq-022", "rfq-025", "rfq-028", "rfq-029"]

READER_FLAGS = {  # flags only a reader of the document can know; the spec states them
    FlagCode.NO_MATCH,
    FlagCode.GRADE_NOT_SPECIFIED,
    FlagCode.UNIT_CONVERSION,
    FlagCode.AMBIGUOUS_UNIT,
    FlagCode.QUANTITY_MISSING,
    FlagCode.INCOTERM_CONFLICT,
    FlagCode.CUSTOMER_UNKNOWN,
}


class _M(BaseModel):
    model_config = ConfigDict(extra="forbid")


class GoldLine(_M):
    mention: str
    cas: str | None
    grade: str | None
    quantity: Decimal | None
    unit: str | None
    packaging: str | None
    sku: str | None
    flags: list[FlagCode] = []
    expect: list[FlagCode] = []
    why: str | None = None


class GoldHeader(_M):
    customer_id: str | None
    incoterms_mentioned: list[str]
    incoterm: str
    destination: str
    currency: str
    requested_date: date | None
    quote_flags: list[FlagCode] = []
    lines: list[GoldLine]


class Document(_M):
    title: str
    fields: dict[str, str]
    columns: list[str]
    rows: list[list[str]]
    notes: str | None = None

    @field_validator("fields", mode="before")
    @classmethod
    def _text_fields(cls, v: dict) -> dict:  # YAML reads 2026-12-04 as a date; a document holds text
        return {str(k): str(x) for k, x in v.items()}

    @field_validator("columns", "rows", mode="before")
    @classmethod
    def _text_cells(cls, v: list) -> list:
        return [[str(c) for c in r] if isinstance(r, list) else str(r) for r in v]


class Spec(_M):
    id: str
    format: str
    received: date
    from_: str | None = None
    subject: str | None = None
    body: str | None = None
    document: Document | None = None
    gold: GoldHeader

    @classmethod
    def load(cls, path: Path) -> Spec:
        raw = yaml.safe_load(path.read_text())
        raw["from_"] = raw.pop("from", None)
        return cls.model_validate(raw)


# ------------------------------------------------------------------------------------------------ rendering


def render_eml(s: Spec, path: Path) -> None:
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = s.from_, SALES_INBOX, s.subject
    msg["Date"] = format_datetime(datetime.combine(s.received, time(9, 30), tzinfo=IST))
    msg["Message-ID"] = f"<{s.id}@quotient-demo.example>"
    msg.set_content(s.body)
    path.write_bytes(bytes(msg))


def _cell(v: str):
    """Numbers typed into Excel are numbers (a CAS typed without dashes becomes one too)."""
    raw = v.replace(",", "")
    try:
        return int(raw)
    except ValueError:
        try:
            return float(raw)
        except ValueError:
            return v


def render_xlsx(s: Spec, path: Path) -> None:
    d = s.document
    wb = Workbook()
    ws = wb.active
    ws.title = "RFQ"
    ws["A1"] = d.title
    ws["A1"].font = Font(bold=True, size=13)
    r = 3
    for k, v in d.fields.items():
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1
    r += 1
    for c, name in enumerate(d.columns, start=1):
        ws.cell(r, c, name).font = Font(bold=True)
    for row in d.rows:
        r += 1
        for c, v in enumerate(row, start=1):
            if v != "":
                ws.cell(r, c, _cell(v))
    for col, width in zip("ABCDEFG", (16, 30, 16, 14, 12, 14, 18), strict=False):
        ws.column_dimensions[col].width = width
    wb.properties.creator = "buyer"
    wb.properties.created = datetime.combine(s.received, time(9, 0))
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(_frozen_zip(buf.getvalue(), wb.properties.created))


def _frozen_zip(data: bytes, when: datetime) -> bytes:
    """openpyxl stamps the save time into docProps and every zip entry; pin both so a rebuild is byte-identical."""
    stamp = when.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp, body)
            dst.writestr(zipfile.ZipInfo(info.filename, date_time=when.timetuple()[:6]), body, zipfile.ZIP_DEFLATED)
    return out.getvalue()


def render_pdf(s: Spec, path: Path) -> None:
    d = s.document
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        title=d.title,
        author="buyer",
        creator="buyer",
        invariant=True,  # no creation timestamp or random document id: rebuilds are byte-identical
    )
    fields = Table(
        [[Paragraph(f"<b>{k}</b>", styles["Normal"]), Paragraph(v, styles["Normal"])] for k, v in d.fields.items()],
        colWidths=[45 * mm, 120 * mm],
    )
    items = Table([d.columns, *d.rows], repeatRows=1)
    items.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    story = [Paragraph(d.title, styles["Title"]), fields, Spacer(1, 8 * mm), items]
    if d.notes:
        story += [Spacer(1, 6 * mm), Paragraph(f"<i>{d.notes}</i>", styles["Normal"])]
    doc.build(story)


RENDER = {"eml": render_eml, "txt": lambda s, p: p.write_text(s.body), "pdf": render_pdf, "xlsx": render_xlsx}


# ------------------------------------------------------------------------------------------------ gold


def _s(d: Decimal | None) -> str | None:
    return None if d is None else str(d)


def gold_for(s: Spec, catalog: Catalog, engine: PricingEngine) -> dict:
    g = s.gold
    country = g.destination
    out_lines, priced = [], []
    for i, ln in enumerate(g.lines, start=1):
        where = f"{s.id} line {i} ({ln.mention})"
        flags = set(ln.flags)
        assert flags <= READER_FLAGS, f"{where}: only reader flags belong in a spec, engine flags go in `expect`"
        row = {
            "line_no": i,
            "mention": ln.mention,
            "cas": ln.cas,
            "grade": ln.grade,
            "quantity": _s(ln.quantity),
            "unit": ln.unit,
            "packaging": ln.packaging,
            "sku": ln.sku,
            "why": ln.why,
        }
        expected = {
            "sku": ln.sku,
            "quantity_uom": None,
            "uom": None,
            "pack_code": None,
            "packs": None,
            "quoted_qty": None,
            "unit_price": None,
            "line_total": None,
            "currency": g.currency,
            "price_row_id": None,
            "freight_row_id": None,
        }

        if ln.sku is None:
            assert FlagCode.NO_MATCH in flags, f"{where}: a line without a SKU must be flagged no_match"
        else:
            p = catalog.product(ln.sku)
            offered = [q.grade for q in catalog.products.values() if q.chemical_code == p.chemical_code]
            if ln.grade is not None:
                assert p.grade == ln.grade, f"{where}: stated grade {ln.grade} but SKU is {p.grade}"
            elif len(offered) > 1:
                assert FlagCode.GRADE_NOT_SPECIFIED in flags, (
                    f"{where}: {p.chemical_name} has grades {offered}; flag it"
                )
                assert p.grade == default_grade(offered), (
                    f"{where}: an unstated grade defaults to {default_grade(offered)}"
                )
            else:
                assert FlagCode.GRADE_NOT_SPECIFIED not in flags, f"{where}: only one grade is offered"
            expected["uom"] = p.uom

            conv = None
            if ln.quantity is None:
                assert FlagCode.QUANTITY_MISSING in flags, f"{where}: no quantity must be flagged"
            else:
                conv = to_uom(ln.quantity, ln.unit, p, catalog.packs, country)
                conv_codes = {f.code for f in conv.flags}
                assert conv_codes == flags & {FlagCode.UNIT_CONVERSION, FlagCode.AMBIGUOUS_UNIT}, (
                    f"{where}: conversion flags {conv_codes} differ from the spec's {flags}"
                )
            if conv is not None and conv.quantity is not None:
                pack = conv.pack_code or ln.packaging
                if conv.pack_code and ln.packaging:
                    assert conv.pack_code == ln.packaging, (
                        f"{where}: unit pack {conv.pack_code} vs packaging {ln.packaging}"
                    )
                line = engine.price(
                    PricingRequest(
                        sku=ln.sku,
                        quantity=conv.quantity,
                        pack_code=pack,
                        customer_id=g.customer_id,
                        incoterm=g.incoterm,
                        destination_country_code=g.destination,
                        currency=g.currency,
                        quote_date=s.received,
                    )
                )
                engine_codes = {f.code for f in line.flags}
                missing = set(ln.expect) - engine_codes
                assert not missing, f"{where}: expected engine flags {missing} did not occur (got {engine_codes})"
                flags |= engine_codes
                expected.update(
                    quantity_uom=_s(conv.quantity),
                    pack_code=line.pack_code,
                    packs=line.packs,
                    quoted_qty=_s(line.quoted_qty),
                    unit_price=_s(line.unit_price),
                    line_total=_s(line.line_total),
                    currency=line.currency,
                    price_row_id=line.price_row_id,
                    freight_row_id=line.freight_row_id,
                )
                priced.append(line)
        severities = {FlagCode(f): None for f in flags}
        row["expected"] = expected
        row["flags"] = sorted(str(f) for f in severities)
        row["needs_human"] = any(SEVERITY[f] == Severity.BLOCK for f in flags) or bool(
            {SEVERITY[f] for f in g.quote_flags} & {Severity.BLOCK}
        )
        out_lines.append(row)

    total = quote_total(priced)
    return {
        "id": s.id,
        "file": f"{s.id}.{s.format}",
        "format": s.format,
        "received": str(s.received),
        "customer_id": g.customer_id,
        "incoterms_mentioned": g.incoterms_mentioned,
        "incoterm": g.incoterm,
        "destination": g.destination,
        "currency": g.currency,
        "requested_date": None if g.requested_date is None else str(g.requested_date),
        "quote_flags": sorted(str(f) for f in g.quote_flags),
        "total": _s(total),
        "lines": out_lines,
        "computed_by": {
            "engine": "quotient.pricing.engine",
            "rules_version": engine.rules.version,
            "data_manifest": "data/manifest.json",
        },
    }


def main() -> None:
    catalog = default_catalog()
    engine = PricingEngine(catalog, default_rules())
    specs = [Spec.load(p) for p in sorted(SPECS.glob("rfq-*.yaml"))]
    assert len({s.id for s in specs}) == len(specs) and set(HOLDOUT) <= {s.id for s in specs}
    RFQS.mkdir(exist_ok=True)
    GOLD.mkdir(exist_ok=True)
    counts: dict[str, int] = {}
    for s in specs:
        RENDER[s.format](s, RFQS / f"{s.id}.{s.format}")
        gold = gold_for(s, catalog, engine)
        (GOLD / f"{s.id}.json").write_text(json.dumps(gold, indent=2, ensure_ascii=False) + "\n")
        for ln in gold["lines"]:
            for f in ln["flags"]:
                counts[f] = counts.get(f, 0) + 1
        for f in gold["quote_flags"]:
            counts[f] = counts.get(f, 0) + 1
    split = {
        "holdout": HOLDOUT,
        "dev": sorted(s.id for s in specs if s.id not in HOLDOUT),
        "note": "Holdout chosen before any extraction prompt was written. Run it only to report, never to tune.",
    }
    (HERE / "split.json").write_text(json.dumps(split, indent=2) + "\n")
    lines = sum(len(s.gold.lines) for s in specs)
    print(
        f"{len(specs)} RFQs, {lines} lines ({len(HOLDOUT)} holdout). Formats: "
        + ", ".join(f"{f} {sum(s.format == f for s in specs)}" for f in RENDER)
    )
    print("flags in gold: " + ", ".join(f"{k} {v}" for k, v in sorted(counts.items(), key=lambda kv: -kv[1])))
    assert GRADE_ORDER


if __name__ == "__main__":
    main()
