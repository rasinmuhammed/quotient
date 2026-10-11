"""Stage 1, ingest: an email, PDF, spreadsheet or pasted text becomes one RFQDocument.

The document's `text` is what the extractor reads and what every extracted span is checked against, so it is
built once here and never rewritten later. Spreadsheet rows keep their columns (" | " between cells) so a
header such as "Qty (MT)" stays readable next to its numbers.
"""

from __future__ import annotations

import email
import io
import re
from datetime import date
from email import policy
from email.utils import parseaddr, parsedate_to_datetime
from pathlib import Path

from quotient.pipeline.types import RFQDocument

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
OWN_DOMAINS = ("quotient-demo.example",)


class IngestError(ValueError):
    pass


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return str(v).strip()


def xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        width = max(max(i for i, c in enumerate(r) if c) + 1 for r in rows)
        if len(wb.worksheets) > 1:
            out.append(f"[Sheet: {ws.title}]")
        out += [" | ".join(r[:width]).rstrip(" |") if any(r[1:width]) else r[0] for r in rows]
    return "\n".join(out)


def pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [p.extract_text() or "" for p in pdf.pages]
    text = "\n\n".join(p.strip() for p in pages if p.strip())
    if not text:
        raise IngestError("the PDF has no text layer (a scan); OCR is not supported in this prototype")
    return text


def _html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", "", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>", "\n", html)
    html = re.sub(r"<[^>]+>", " ", html)
    return re.sub(r"[ \t]+", " ", html).strip()


def eml_document(data: bytes, name: str) -> RFQDocument:
    msg = email.message_from_bytes(data, policy=policy.default)
    sender_name, sender_email = parseaddr(str(msg.get("From", "")))
    received = None
    if msg.get("Date"):
        try:
            received = parsedate_to_datetime(str(msg["Date"])).date()
        except (TypeError, ValueError):
            received = None
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if body_part is not None:
        body = body_part.get_content()
        if body_part.get_content_type() == "text/html":
            body = _html_to_text(body)
    parts = [f"From: {msg.get('From', '')}", f"Subject: {msg.get('Subject', '')}", f"Date: {msg.get('Date', '')}", "",
             body.strip()]
    for att in msg.iter_attachments():
        fname = att.get_filename() or "attachment"
        payload = att.get_content()
        payload = payload if isinstance(payload, bytes) else str(payload).encode()
        ext = Path(fname).suffix.lower()
        if ext == ".pdf":
            parts += ["", f"--- Attachment: {fname} ---", pdf_text(payload)]
        elif ext in (".xlsx", ".xlsm"):
            parts += ["", f"--- Attachment: {fname} ---", xlsx_text(payload)]
        elif ext in (".txt", ".csv"):
            parts += ["", f"--- Attachment: {fname} ---", payload.decode("utf-8", "replace")]
    return RFQDocument(source_name=name, format="eml", text="\n".join(parts).strip(), sender_email=sender_email or None,
                       sender_name=sender_name or None, subject=str(msg.get("Subject", "")) or None, received=received)


def _sender_from_text(text: str) -> str | None:
    found = [e for e in EMAIL_RE.findall(text) if not e.lower().endswith(OWN_DOMAINS)]
    return found[0] if found else None


def ingest_text(text: str, name: str = "pasted.txt", received: date | None = None) -> RFQDocument:
    text = text.replace("\r\n", "\n").strip()
    if not text:
        raise IngestError("the RFQ is empty")
    return RFQDocument(source_name=name, format="text", text=text, sender_email=_sender_from_text(text), received=received)


def ingest(data: bytes, name: str) -> RFQDocument:
    ext = Path(name).suffix.lower()
    if ext == ".eml":
        return eml_document(data, name)
    if ext == ".pdf":
        text = pdf_text(data)
        return RFQDocument(source_name=name, format="pdf", text=text, sender_email=_sender_from_text(text))
    if ext in (".xlsx", ".xlsm"):
        text = xlsx_text(data)
        return RFQDocument(source_name=name, format="xlsx", text=text, sender_email=_sender_from_text(text))
    if ext in (".txt", ".text", ""):
        try:
            decoded = data.decode("utf-8")
        except UnicodeDecodeError:
            decoded = data.decode("latin-1")
        return ingest_text(decoded, name)
    raise IngestError(f"unsupported file type {ext!r}: send an email (.eml), PDF, Excel (.xlsx) or text")
